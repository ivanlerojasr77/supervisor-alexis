import json
import os
import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from starlette.datastructures import UploadFile

import run_agent


class ReportAnalysisTests(unittest.TestCase):
    def setUp(self):
        state = SimpleNamespace(
            credential=object(),
            endpoint="https://example.services.ai.azure.com/api/projects/test",
            agent_names={
                "supervisor": "supervisor-alexis",
                "financial": "analista-financiero",
                "technical": "analista-tecnico",
            },
        )
        self.request = SimpleNamespace(app=SimpleNamespace(state=state))

    def upload(self, filename: str, content: bytes) -> UploadFile:
        return UploadFile(filename=filename, file=BytesIO(content))

    def test_builds_managed_voice_websocket_url(self):
        url = run_agent.build_voice_url(
            "https://example.services.ai.azure.com/api/projects/test",
            "supervisor-alexis",
        )

        self.assertTrue(url.startswith("wss://"))
        self.assertIn(
            "/agents/supervisor-alexis/endpoint/protocols/voice?api-version=v1",
            url,
        )
        self.assertIn("agent_session_id=", url)

    def test_extracts_text_from_txt(self):
        self.assertEqual(
            run_agent.extract_report_text("report.txt", b"\xef\xbb\xbfIngresos: 200"),
            "Ingresos: 200",
        )

    def test_extracts_paragraphs_and_tables_from_docx(self):
        document = Document()
        document.add_paragraph("Resumen financiero")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Ingresos"
        table.cell(0, 1).text = "200"
        buffer = BytesIO()
        document.save(buffer)

        text = run_agent.extract_report_text("report.docx", buffer.getvalue())

        self.assertIn("Resumen financiero", text)
        self.assertIn("Ingresos | 200", text)

    def test_routes_financial_classification_to_financial_analyst(self):
        classification = {
            "category": "financial",
            "rationale": "El reporte presenta ingresos y costos.",
            "confidence": 0.94,
        }
        with patch.object(
            run_agent,
            "invoke_agent",
            side_effect=[json.dumps(classification), "Ingresos: 200; costos: 80."],
        ) as invoke:
            result = run_agent.analyze_report(
                self.request,
                self.upload("reporte.txt", b"Ingresos: 200; costos: 80."),
                question="¿Cuáles fueron las variaciones?",
            )

        self.assertEqual(result["classification"], classification)
        self.assertEqual(result["analysis"], "Ingresos: 200; costos: 80.")
        self.assertEqual(invoke.call_args_list[0].args[2], "supervisor-alexis")
        self.assertNotIn("variaciones", invoke.call_args_list[0].args[3])
        self.assertEqual(invoke.call_args_list[1].args[2], "analista-financiero")
        self.assertIn("¿Cuáles fueron las variaciones?", invoke.call_args_list[1].args[3])

    def test_routes_technical_classification_to_technical_analyst(self):
        classification = {
            "category": "technical",
            "rationale": "El reporte describe una interrupción de servicio.",
            "confidence": 0.9,
        }
        with patch.object(
            run_agent,
            "invoke_agent",
            side_effect=[json.dumps(classification), "Disponibilidad: 96 %."],
        ) as invoke:
            result = run_agent.analyze_report(
                self.request,
                self.upload("reporte.txt", b"Interrupcion de servicio"),
                question="",
            )

        self.assertEqual(result["analysis"], "Disponibilidad: 96 %.")
        self.assertEqual(invoke.call_args_list[1].args[2], "analista-tecnico")

    def test_uncertain_classification_does_not_call_an_analyst(self):
        classification = {
            "category": "uncertain",
            "rationale": "El reporte mezcla temas sin evidencia suficiente.",
            "confidence": 0.4,
        }
        with patch.object(
            run_agent,
            "invoke_agent",
            return_value=json.dumps(classification),
        ) as invoke:
            result = run_agent.analyze_report(
                self.request,
                self.upload("reporte.txt", b"Reporte mixto"),
                question="",
            )

        self.assertIsNone(result["analysis"])
        invoke.assert_called_once()

    def test_rejects_unsupported_file_format(self):
        with self.assertRaises(HTTPException) as error:
            run_agent.analyze_report(
                self.request,
                self.upload("reporte.csv", b"ingresos,200"),
                question="",
            )

        self.assertEqual(error.exception.status_code, 415)

    def test_rejects_invalid_supervisor_json(self):
        with patch.object(run_agent, "invoke_agent", return_value="Clasificación financiera"):
            with self.assertRaises(HTTPException) as error:
                run_agent.analyze_report(
                    self.request,
                    self.upload("reporte.txt", b"Ingresos: 200"),
                    question="",
                )

        self.assertEqual(error.exception.status_code, 502)

    def test_normalizes_voice_transcript_wrappers(self):
        classification = {
            "category": "financial",
            "rationale": "Predominan ingresos y costos.",
            "confidence": 0.9,
        }
        with patch.object(
            run_agent,
            "invoke_agent",
            side_effect=[f"=({json.dumps(classification)})>", "Resumen financiero."],
        ):
            result = run_agent.analyze_report(
                self.request,
                self.upload("reporte.txt", b"Ingresos: 200; costos: 80."),
                question="",
            )

        self.assertEqual(result["classification"], classification)

    def test_rejects_extra_supervisor_json_properties(self):
        output = json.dumps(
            {
                "category": "financial",
                "rationale": "Predominan ingresos y costos.",
                "confidence": 0.9,
                "analysis": "No debe incluir análisis aquí.",
            }
        )
        with patch.object(run_agent, "invoke_agent", return_value=output):
            with self.assertRaises(HTTPException) as error:
                run_agent.analyze_report(
                    self.request,
                    self.upload("reporte.txt", b"Ingresos: 200; costos: 80."),
                    question="",
                )

        self.assertEqual(error.exception.status_code, 502)

    def test_http_upload_parses_question_and_returns_analysis(self):
        credential = MagicMock()
        credential.__enter__.return_value = credential
        classification = {
            "category": "financial",
            "rationale": "El reporte presenta ingresos.",
            "confidence": 0.9,
        }
        with (
            patch.dict(
                os.environ,
                {
                    "AZURE_AI_PROJECT_ENDPOINT": "https://example.services.ai.azure.com/api/projects/test",
                    "SUPERVISOR_AGENT_NAME": "supervisor-alexis",
                    "FINANCIAL_AGENT_NAME": "analista-financiero",
                    "TECHNICAL_AGENT_NAME": "analista-tecnico",
                },
            ),
            patch.object(run_agent, "DefaultAzureCredential", return_value=credential),
            patch.object(
                run_agent,
                "invoke_agent",
                side_effect=[json.dumps(classification), "Ingresos: 200."],
            ) as invoke,
        ):
            with TestClient(run_agent.app) as client:
                response = client.post(
                    "/api/reports/analyze",
                    files={"file": ("reporte.txt", b"Ingresos: 200.", "text/plain")},
                    data={"question": "Resume los ingresos."},
                )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["analysis"], "Ingresos: 200.")
        self.assertIn("Resume los ingresos.", invoke.call_args_list[1].args[3])


if __name__ == "__main__":
    unittest.main()
